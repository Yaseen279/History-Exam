"""Build the graduation research paper (.docx) from the analysis outputs.

Usage: python3 04_build_docx.py <outdir> <docx path>
College format: Times New Roman 14; Abstract (<=1.5 pages); Introduction (>=4 pages);
Method; Results; Discussion; Conclusions and Recommendations; References (Vancouver).
Preliminary pages (Title, Acknowledgment, Abbreviations, Tables, Figures, Abstract) use
roman numerals; arabic numbering starts at 1 on the Introduction.
"""
import sys, json, re
import numpy as np, pandas as pd
from scipy import stats
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from refs import REFS

OUT, DOCX = sys.argv[1], sys.argv[2]
d = pd.read_pickle(f'{OUT}/analysed.pkl')
X = pd.read_excel(f'{OUT}/tables/all_results.xlsx', sheet_name=None, index_col=0)
R = json.load(open(f'{OUT}/summary.json'))
n = len(d)
FONT = 'Times New Roman'

# ------------------------------------------------------------------ helpers
def f2(x): return f'{x:.2f}'
def f1(x): return f'{x:.1f}'
def f3(x): return f'{x:.3f}'
def pv(p):
    return 'p < 0.001' if p < 0.001 else f'p = {p:.3f}'
def pnum(p):
    return '<0.001' if p < 0.001 else f'{p:.3f}'
def pct(k, tot=n): return f'{k/tot*100:.1f}%'
def npct(k, tot=n): return f'{k} ({k/tot*100:.1f}%)'
def msd(s): return f'{s.mean():.2f} ± {s.std():.2f}'

order = []
def cite(*keys):
    nums = []
    for k in keys:
        assert k in REFS, k
        if k not in order: order.append(k)
        nums.append(order.index(k) + 1)
    nums = sorted(set(nums))
    # compress ranges: 1,2,3 -> 1-3
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j+1] == nums[j] + 1: j += 1
        out.append(f'{nums[i]}–{nums[j]}' if j - i >= 2 else (f'{nums[i]},{nums[j]}' if j == i + 1 else f'{nums[i]}'))
        i = j + 1
    return '[' + ','.join(out) + ']'

doc = Document()
st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(14)
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT); st.element.rPr.rFonts.set(qn('w:cs'), FONT)
pf = st.paragraph_format; pf.line_spacing = 1.5; pf.space_after = Pt(6)
for lvl, size in [(1, 16), (2, 14), (3, 14)]:
    h = doc.styles[f'Heading {lvl}']; h.font.name = FONT; h.font.size = Pt(size); h.font.bold = True
    h.font.italic = (lvl == 3); h.font.color.rgb = RGBColor(0, 0, 0)
    rpr = h.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
    if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
    for a in list(rf.attrib):
        if a.endswith('Theme'): del rf.attrib[a]   # theme fonts would override Times New Roman
    for a in ['w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs']: rf.set(qn(a), FONT)
    h.paragraph_format.space_before = Pt(12 if lvl > 1 else 0); h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.line_spacing = 1.5; h.paragraph_format.keep_with_next = True
sec = doc.sections[0]
for s in [sec]:
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = Cm(3.0); s.right_margin = Cm(2.5); s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)

def add_runs(p, text, size=None):
    # **bold**, *italic*
    parts = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('*'): r = p.add_run(part[1:-1]); r.italic = True
        else: r = p.add_run(part)
        r.font.name = FONT
        if size: r.font.size = Pt(size)
    return p
def P(text, align='justify', size=None, indent=True, space_after=6, bold=False):
    p = doc.add_paragraph()
    p.alignment = {'justify': WD_ALIGN_PARAGRAPH.JUSTIFY, 'center': WD_ALIGN_PARAGRAPH.CENTER, 'left': WD_ALIGN_PARAGRAPH.LEFT}[align]
    if indent and align == 'justify': p.paragraph_format.first_line_indent = Cm(1.0)
    p.paragraph_format.space_after = Pt(space_after)
    add_runs(p, f'**{text}**' if bold else text, size)
    return p
def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); add_runs(p, text); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3); return p
def numbered(text):
    p = doc.add_paragraph(style='List Number'); add_runs(p, text); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(3); return p
HEADS = []
def H(text, lvl=1):
    HEADS.append((lvl, text)); return doc.add_heading(text, level=lvl)
def page_break(): doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def add_field(run, instr, cached='1'):
    f1_ = OxmlElement('w:fldChar'); f1_.set(qn('w:fldCharType'), 'begin')
    it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr
    f2_ = OxmlElement('w:fldChar'); f2_.set(qn('w:fldCharType'), 'separate')
    t = OxmlElement('w:t'); t.text = cached
    f3_ = OxmlElement('w:fldChar'); f3_.set(qn('w:fldCharType'), 'end')
    for e in (f1_, it, f2_, t, f3_): run._r.append(e)

tables_list, figures_list = [], []
def caption(kind, text):
    lst = tables_list if kind == 'Table' else figures_list
    lst.append(text); num = len(lst)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.keep_with_next = (kind == 'Table')
    r = p.add_run(f'{kind} '); r.bold = True; r.font.size = Pt(12)
    # field split over separate runs so the cached number keeps the caption formatting
    for part in ['begin', 'instr', 'separate', 'text', 'end']:
        r = p.add_run(); r.bold = True; r.font.size = Pt(12)
        if part == 'instr':
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f' SEQ {kind} \\* ARABIC '; r._r.append(it)
        elif part == 'text':
            r.text = str(num)
        else:
            e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), part); r._r.append(e)
    r = p.add_run('. '); r.bold = True; r.font.size = Pt(12)
    r2 = p.add_run(text); r2.font.size = Pt(12); r2.italic = (kind == 'Figure')
    return num

def set_cell_border(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr(); b = tcPr.find(qn('w:tcBorders'))
    if b is None: b = OxmlElement('w:tcBorders'); tcPr.append(b)
    for edge, val in kw.items():
        el = OxmlElement(f'w:{edge}'); el.set(qn('w:val'), val); el.set(qn('w:sz'), '8'); el.set(qn('w:color'), '000000'); b.append(el)

def table(header, rows, widths=None, size=10.5, note=None, bold_rows=()):
    t = doc.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = doc.styles['Table Grid']
    # three-line (APA-like) table: remove inner borders
    tblPr = t._tbl.tblPr; bd = tblPr.find(qn('w:tblBorders'))
    if bd is not None: tblPr.remove(bd)
    def fill(cells, vals, bold=False, italic_first=False):
        for c, v in zip(cells, vals):
            c.text = ''; p = c.paragraphs[0]; p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(0)
            v = '0.00' if str(v) == '-0.00' else v
            r = p.add_run(str(v)); r.font.size = Pt(size); r.bold = bold; r.font.name = FONT
    fill(t.rows[0].cells, header, bold=True)
    for c in t.rows[0].cells: set_cell_border(c, top='single', bottom='single')
    for i, row in enumerate(rows):
        cells = t.add_row().cells; fill(cells, row, bold=(i in bold_rows))
    for c in t.rows[-1].cells: set_cell_border(c, bottom='single')
    for row in t.rows:
        tr = row._tr; trPr = tr.get_or_add_trPr(); cs = OxmlElement('w:cantSplit'); trPr.append(cs)
    if widths:
        t.autofit = False
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); t._tbl.tblPr.append(lay)
        for col, w in zip(t.columns, widths): col.width = Cm(w)
        for row in t.rows:
            for c, w in zip(row.cells, widths): c.width = Cm(w)
    if note:
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(10)
        add_runs(p, note, size=9.5)
    else:
        doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

def figure(path, text, width=15.5):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width))
    caption('Figure', text)

FIG = f'{OUT}/figures'

# ------------------------------------------------------------------ numbers used in text
dates = R['dates']
mc = R['main_corr']; sc = mc['SHAI_CSS']
part = R['partial_corr']
prev = R['prevalences']; ss = R['sample_size']
rel = X['T4_reliability'].set_index('Scale')
cfa = X['T6_cfa']
hier = X['T13_hierarchical_CSS']
ols = X['T14_OLS_CSS_full']; olsS = X['T15_OLS_SHAI_full']
lg = X['T18_logit_HA_clin']; beh = X['T19_logit_behaviours']; ordl = X['T20_ordinal_logit']
med = X['T21_mediation']; ser = X['T22_serial_mediation']; mod = X['T23_moderation']
sg = X['T26_subgroup_corr']; lang = X['T24_language']; ms = X['T25_medstudents']
G = X['T9_group_comparisons']; CH = X['T10_chi_square']; cmpc = X['T7b_compare_corr']
sens = R['sensitivity']; quant = X['T17_quantile_reg']; subq = X['T16_subscales_on_SHAI']
def gp(var, out): return G[(G.Variable == var) & (G.Outcome == out)].iloc[0]
def chi(out, var): return CH[(CH.Outcome.str.startswith(out)) & (CH.Variable == var)].iloc[0]
def bh(outcome, pred): return beh[(beh.Outcome == outcome) & (beh.Predictor == pred)].iloc[0]
def medrow(model_prefix, path): return med[med.Model.str.startswith(model_prefix) & (med.Path == path)].iloc[0]
def ci(lo, hi, k=2): return f'{lo:.{k}f}–{hi:.{k}f}'

shai, css = d.SHAI, d.CSS
ib = R['item_by_item']
import statsmodels.formula.api as smf
_m = smf.ols("CSS ~ C(age_c, Treatment('20–29')) + female + C(occ_c, Treatment('Non-health employed/other')) + chronic + lang_en + tool_ai + tool_search + tool_medplat + tool_social + trust + mot_reassurance + mot_selfdx + mot_cost + mot_triage + SHAI", d).fit(cov_type='HC3')
ai_alt_B, ai_alt_p = _m.params['tool_ai'], _m.pvalues['tool_ai']
kw_bot = stats.kruskal(*[d.loc[d.ai_model == m, 'CSS'] for m in ['ChatGPT', 'Gemini', 'Claude']]).pvalue
_df2 = n - 20 - 1; pow20 = float(1 - stats.ncf.cdf(stats.f.ppf(0.95, 1, _df2), 1, _df2, 0.03 * n))
cnt = lambda c: int(d[c].sum())
age20 = int((d.age_group == '20-29').sum()); fem = cnt('female'); medst = cnt('medical_student')
ai_n = cnt('ai_user'); se_n = cnt('tool_search'); hb = cnt('health_background')
r_sub = {k: mc[f'SHAI_{k}'] for k in ['Excessiveness', 'Compulsion', 'Distress', 'Reassurance']}
mA_ind = medrow('A:', 'Indirect'); mA_dir = medrow('A:', 'Direct'); mA_tot = medrow('A:', 'Total')
mC_ind = medrow('C:', 'Indirect'); mC_dir = medrow('C:', 'Direct'); mC_tot = medrow('C:', 'Total')
mD_ind = medrow('D:', 'Indirect'); mB_ind = medrow('B:', 'Indirect')
pmed = R['prop_mediated']
b_shai = ols.loc['SHAI']; b_ai = ols.loc['ai_user']
lg_css = lg.loc['zCSS']; lg_mot = lg.loc['mot_reassurance']; lg_age = lg.loc['age_ord']
hiS = olsS.loc['CSS']; hiChron = olsS.loc['chronic']; hiMot = olsS.loc['mot_reassurance']

# ================================================================== TITLE PAGE (section 1: roman)
P('', indent=False)
P('[College logo]', align='center', indent=False)
P('[University name]', align='center', indent=False, bold=True)
P('[College name]', align='center', indent=False, bold=True)
P('', indent=False)
p = P('Artificial Intelligence, Search Engines, and Cyberchondria: A Study of Patient Pre-Consultation Behavior in Baghdad',
      align='center', indent=False, size=18, bold=True)
P('', indent=False)
P('A graduation research project submitted in partial fulfilment of the requirements for the degree of [Degree]', align='center', indent=False)
P('', indent=False)
P('**Prepared by:**', align='center', indent=False)
P('[Student name 1]   ·   [Student name 2]   ·   [Student name 3]', align='center', indent=False)
P('', indent=False)
P('**Supervised by:**', align='center', indent=False)
P('[Supervisor name and title]', align='center', indent=False)
P('', indent=False)
P('[Academic year]', align='center', indent=False)
page_break()

# ACKNOWLEDGMENT
H('Acknowledgment')
P('[To be completed by the authors.]', indent=False)
P('We thank all the participants who gave their time to complete this survey, and [names of supervisor, department and any others to be acknowledged].')
page_break()

# TABLE OF CONTENTS (filled at the end)
H('Table of Contents'); TOC_anchor = doc.add_paragraph(); page_break()

# LIST OF ABBREVIATIONS
H('List of Abbreviations')
abbr = [('AI', 'Artificial intelligence'), ('AOR', 'Adjusted odds ratio'), ('AUC', 'Area under the receiver operating characteristic curve'),
        ('CFA', 'Confirmatory factor analysis'), ('CFI', 'Comparative fit index'), ('CI', 'Confidence interval'),
        ('CSS-12', 'Cyberchondria Severity Scale, 12-item short form'), ('DWLS', 'Diagonally weighted least squares'),
        ('ER', 'Emergency room'), ('FDR', 'False discovery rate'), ('GP', 'General practitioner'), ('HA', 'Health anxiety'),
        ('HC3', 'Heteroskedasticity-consistent standard errors, type 3'), ('IQR', 'Interquartile range'),
        ('KMO', 'Kaiser–Meyer–Olkin measure of sampling adequacy'), ('LLM', 'Large language model'), ('ML', 'Maximum likelihood'),
        ('OLS', 'Ordinary least squares'), ('OR', 'Odds ratio'), ('RMSEA', 'Root mean square error of approximation'),
        ('SD', 'Standard deviation'), ('SHAI-14', 'Short Health Anxiety Inventory, 14-item version'),
        ('STROBE', 'Strengthening the Reporting of Observational Studies in Epidemiology'), ('TLI', 'Tucker–Lewis index'),
        ('VIF', 'Variance inflation factor')]
t = doc.add_table(rows=0, cols=2)
for a, b in abbr:
    c = t.add_row().cells; c[0].text = a; c[1].text = b
    for cc in c:
        for r in cc.paragraphs[0].runs: r.font.size = Pt(13)
        cc.paragraphs[0].paragraph_format.space_after = Pt(1); cc.paragraphs[0].paragraph_format.line_spacing = 1.15
    c[0].paragraphs[0].runs[0].bold = True
for row in t.rows: row.cells[0].width = Cm(3.5); row.cells[1].width = Cm(12)
page_break()

# placeholders for lists of tables / figures (filled at the end)
H('List of Tables'); LT_anchor = doc.add_paragraph(); page_break()
H('List of Figures'); LF_anchor = doc.add_paragraph(); page_break()

# ================================================================== ABSTRACT
H('Summary / Abstract')
sd_ = lambda s: f'{s.mean():.1f} ± {s.std():.1f}'
abstract = [
 f'**Background:** Patients increasingly consult search engines and, more recently, artificial-intelligence (AI) chatbots before seeing a doctor. Health anxiety and cyberchondria (excessive, distressing online health searching) are thought to reinforce each other, but evidence from Iraq, and on AI chatbots, is scarce.',
 f'**Objectives:** To measure the association between health anxiety and cyberchondria among adults in Baghdad, to identify the socio-demographic and digital-behaviour factors linked with each, and to describe how online searching shapes pre-consultation behaviour and the doctor–patient relationship.',
 f'**Methods:** A cross-sectional online survey (Google Forms) of {n} adults in Baghdad was conducted between {pd.Timestamp(dates[0]):%d %B %Y} and {pd.Timestamp(dates[1]):%d %B %Y}. Health anxiety was measured with the Short Health Anxiety Inventory (SHAI-14) and cyberchondria with the Cyberchondria Severity Scale short form (CSS-12). Analyses included reliability and confirmatory factor analysis, Spearman and partial correlations, non-parametric group comparisons with false-discovery-rate correction, hierarchical robust regression, logistic and ordinal regression, and bootstrapped mediation and moderation models.',
 f'**Results:** Participants were mostly young (aged 20–29: {pct(age20)}), female ({pct(fem)}) and from health-related fields ({pct(hb)}). AI chatbots were the most used tool ({pct(ai_n)}), ahead of search engines ({pct(se_n)}). Both scales were reliable (α = {rel.loc["SHAI-14","alpha"]:.2f} and {rel.loc["CSS-12","alpha"]:.2f}). Mean SHAI was {sd_(shai)} and {prev["HA_clin"]:.1f}% scored ≥18. Health anxiety correlated moderately with cyberchondria (r = {sc["pearson_r"]:.2f}, 95% CI {ci(*sc["pearson_CI"])}), independent of covariates (partial r = {part["r"]:.2f}), most strongly with excessive searching and post-search distress. Health anxiety was the strongest independent predictor of cyberchondria (β = {b_shai.beta_std:.2f}), adding {hier.dR2.iloc[2]*100:.1f}% explained variance; AI-chatbot use was associated with higher cyberchondria (B = {b_ai.B:.1f}), although this depended on how inconsistent answers were classified. Excessive searching partly mediated the link between health anxiety and post-search distress, and post-search distress almost fully mediated its link with reassurance-seeking consultations. After searching, {pct(cnt("beh_delay_reassured"))} delayed or skipped care because they felt reassured, and {pct(cnt("share_never"))} never disclosed their searches to their doctor for fear of being judged.',
 f'**Conclusion:** Health anxiety and cyberchondria travel together in Baghdad, and AI chatbots are now the main gateway to online health information. Clinicians should ask about online and AI searches, and digital health literacy and anxiety-aware AI design are recommended. Longitudinal work is needed to confirm the direction of the cycle.',
 '**Keywords:** cyberchondria; health anxiety; artificial intelligence; chatbots; online health information seeking; Iraq.']
for a in abstract: P(a, space_after=4)

# ================================================================== MAIN TEXT (section 2: arabic from 1)
new = doc.add_section(WD_SECTION.NEW_PAGE)

H('1. Introduction')
H('1.1 Background', 2)
P(f'Searching the Internet before seeing a doctor has become an ordinary part of being ill. A symptom is now often typed into a search engine, or described to an AI chatbot, long before it is described to a physician. This shift has clear benefits: information is immediate, free and available at any hour, and it can help people judge whether a symptom needs medical attention. It also has a well-documented downside. For a sizeable minority of people, online searching does not settle a worry but amplifies it, a phenomenon that has come to be called *cyberchondria* {cite("white2009","starcevic2013")}.')
P(f'Iraq offers a useful and under-studied setting in which to examine this. Internet access in the country has expanded rapidly, with roughly four in five Iraqis online by the start of 2025 {cite("datareportal2025")}, and in such a highly connected population the Internet and, increasingly, AI chatbots may function as a first point of contact for health concerns. How this behaviour interacts with health anxiety, and what it means for the consultation that follows, has not been examined in Baghdad.')

H('1.2 Health anxiety', 2)
P(f'Health anxiety refers to a persistent preoccupation with having, or acquiring, a serious illness, maintained by misinterpretation of normal bodily sensations. It lies on a continuum, from transient worry that most people experience to clinically significant illness anxiety disorder and somatic symptom disorder in the DSM-5 {cite("apa2013")}. The cognitive-behavioural model proposes that bodily sensations are interpreted as signs of serious disease and that the resulting anxiety is maintained by safety behaviours such as body checking and reassurance seeking: these behaviours relieve anxiety briefly but prevent the person from learning that the feared outcome does not occur, so the worry returns {cite("salk1986")}.')
P(f'Health anxiety is not rare. In a national Australian survey, 5.7% of adults met lifetime criteria and 3.4% current criteria for health anxiety {cite("sunderland2013")}, and sub-threshold forms are considerably more common. Its costs are substantial: patients in the top band of health anxiety and somatisation in primary care had markedly more physician visits and higher outpatient costs than other patients, even after adjustment for medical comorbidity {cite("barsky2001")}. In Iraq, almost half of 1591 respondents in an online survey during the COVID-19 restrictions reported health anxiety {cite("karim2020")}, suggesting that the population is far from immune to illness worry.')

H('1.3 Cyberchondria', 2)
P(f'The term cyberchondria was popularised by White and Horvitz, whose analysis of large-scale search logs showed that searches for common, benign symptoms frequently escalated towards rare and serious diagnoses, and that this escalation was driven by the amount and wording of medical content viewed and by the user\'s own predisposition to escalate {cite("white2009")}. Starcevic and Berle later defined cyberchondria as excessive or repeated online searching for health information that is associated with increasing, rather than relieving, distress or anxiety {cite("starcevic2013")}. Crucially, searching for health information is not itself pathological; it becomes cyberchondria when it is excessive, compulsive and distressing, and when it drives further reassurance seeking {cite("starcevic2017","starcevic2020")}. A systematic review described cyberchondria as a possible transdiagnostic, compulsive digital syndrome closely linked to health anxiety and obsessive-compulsive symptoms, while noting that the evidence base was almost entirely cross-sectional and based on online student samples {cite("vismara2020")}.')
P(f'The most widely used instrument is the Cyberchondria Severity Scale (CSS), originally a 33-item measure with five factors: compulsion, distress, excessiveness, reassurance seeking and mistrust of medical professionals {cite("mcelroy2014")}. Its 12-item short form (CSS-12) drops the mistrust dimension and retains four three-item subscales that load substantially on a general cyberchondria factor, supporting both a total and subscale scores {cite("mcelroy2019")}. An Arabic version of the CSS-12 has shown acceptable structure and validity among Lebanese adults {cite("hallit2023")}.')

H('1.4 The relationship between health anxiety and cyberchondria', 2)
P(f'People with higher health anxiety search for health information more often, spend longer doing so and find searching more distressing than people with low health anxiety {cite("muse2012")}. Among people who search online for symptoms, those with high illness anxiety tend to feel worse after searching whereas those with low illness anxiety tend to feel relieved {cite("doherty2016")}. Two meta-analyses have quantified the association: health anxiety correlates moderately with general online health information seeking (r = 0.34) but strongly with cyberchondria (r = 0.62) {cite("mcmullan2019")}, and a later meta-analysis reported a similar pooled correlation (r = 0.63), with subscale correlations ranging from 0.24 to 0.66 {cite("schenkel2021")}. The same work frames health-related Internet use in anxious individuals as a safety behaviour maintained by intermittent reinforcement {cite("schenkel2021")}. Intolerance of uncertainty strengthens the link between searching and anxiety, presumably because search results rarely deliver a single, certain answer {cite("fergus2013")}.')
P(f'Whether the two conditions "fuel each other" is the central question. The only large longitudinal study, a four-wave Dutch panel of 5322 adults, found initial evidence of a reciprocal relationship in which health-anxious people search more and searching in turn raises health anxiety; however, this loop appeared among people with non-clinical health anxiety and not in the clinically anxious subsample, in whom searching seemed to maintain rather than worsen anxiety {cite("tepoel2016")}. The COVID-19 pandemic offered a natural test of the cycle: cyberchondria tracked rising concern and safety behaviours in Croatia {cite("jokic2020")}, and was linked with health anxiety among university students in Turkey {cite("kurcer2022")}. A recent scoping review reported cyberchondria prevalence estimates between 30.7% and 55.6% across studies and consistent associations with health anxiety {cite("miezah2026")}.')

H('1.5 From search engines to AI chatbots', 2)
P(f'Most of this evidence was gathered when "searching online" meant using a search engine. Large language model (LLM) chatbots such as ChatGPT and Gemini change the experience in important ways: they answer in fluent, personalised, conversational language; they can be questioned repeatedly; and their responses can appear more empathetic than those of physicians {cite("thiru2023","ayers2023")}. In a large US survey, about one in five respondents had used ChatGPT for health information, most of them remained sceptical of its output, and most then turned to their physicians {cite("ayo2024")}. Health anxiety has also been identified as a stable predictor of symptom-checker app use, with over half of users meeting a hypochondria classification {cite("wetzel2024")}. Whether AI chatbots relieve worry, by offering calm and contextualised answers, or amplify it, by being endlessly available for another round of reassurance, is an open empirical question.')

H('1.6 Online searching and the doctor–patient relationship', 2)
P(f'What patients do after searching matters for clinical care. Online information can improve the doctor–patient relationship when patients discuss it with their physician, but can strain it when it is withheld or used to challenge the physician {cite("tan2017")}. Reassurance seeking may translate into additional consultations, requests for investigations or "doctor shopping", while false reassurance may lead to delayed care. Studies in the Arab region have begun to describe these patterns: cyberchondria was common among adults in Jeddah and associated with smartphone addiction {cite("elzayat2023")}, and among Saudi primary-care patients cyberchondria was linked with health anxiety {cite("pasayan2025")}. No comparable study has been reported from Iraq, and none, to our knowledge, has placed AI chatbots alongside search engines in this population.')

H('1.7 Rationale', 2)
P('This study addresses three gaps. First, it provides the first estimate of the association between health anxiety and cyberchondria among adults in Baghdad. Second, it examines AI chatbots as a distinct, and now dominant, source of online health information, rather than treating "the Internet" as one entity. Third, it links these psychological constructs to concrete pre-consultation behaviours, namely delaying care, self-medication, urgent visits, multiple consultations, disclosure of searches and trust in physicians, that matter directly to clinicians.')

H('1.8 Aim and objectives', 2)
P('**Aim:** To study the relationship between health anxiety and cyberchondria and their correlates among adults in Baghdad who search for health information online, with particular attention to AI chatbots and pre-consultation behaviour.', indent=False)
P('**Specific objectives:**', indent=False)
for o in ['To describe the tools, motivations and behavioural responses associated with online health information seeking, including the use of AI chatbots.',
          'To estimate the level of health anxiety (SHAI-14) and cyberchondria (CSS-12) and the proportion with clinically relevant health anxiety.',
          'To measure the strength of the association between health anxiety and cyberchondria overall and for each cyberchondria dimension.',
          'To test the association of every socio-demographic, clinical and digital-use variable with health anxiety and cyberchondria.',
          'To identify independent predictors of each construct using multivariable models.',
          'To examine pathways consistent with a mutually reinforcing cycle (mediation) and factors that modify the association (moderation).',
          'To examine the links between health anxiety, cyberchondria and pre-consultation behaviour, disclosure of searches to physicians and trust in physicians.']:
    numbered(o)
P('**Hypotheses:** (H1) Health anxiety is positively associated with cyberchondria, most strongly with its distress and excessiveness dimensions. (H2) The association is independent of socio-demographic and digital-use factors. (H3) Excessive searching statistically mediates the link between health anxiety and post-search distress. (H4) Higher cyberchondria is associated with care-seeking behaviours (urgent visits, multiple consultations) and with a more strained doctor–patient relationship.', indent=False)

# ------------------------------------------------------------------ METHOD
H('2. Method')
H('2.1 Study design and setting', 2)
P(f'This was an analytical cross-sectional study reported in line with the STROBE statement for observational studies {cite("strobe2007")} and the CHERRIES checklist for web-based surveys {cite("eysenbach2004")}. Data were collected online in Baghdad, Iraq, between {pd.Timestamp(dates[0]):%d %B %Y} and {pd.Timestamp(dates[1]):%d %B %Y}.')
H('2.2 Participants and sampling', 2)
P('A non-probability (convenience and snowball) sampling approach was used. The survey link was distributed through online channels (social media and messaging groups) whose members were based in Baghdad. Inclusion criteria were: age 18 years or older, ability to read Arabic or English, and provision of informed consent. Every completed submission was retained for analysis; three responses that the research team had annotated during data collection were kept and examined in a sensitivity analysis.')
H('2.3 Sample size', 2)
P(f'The study was powered for its primary objective, the correlation between health anxiety and cyberchondria. Using Fisher\'s z approximation with two-sided α = 0.05 and 80% power, a minimum of {ss["n_needed_r_0_20"]} participants is required to detect a small-to-moderate correlation (r = 0.20), and {ss["n_needed_r_0_30"]} to detect r = 0.30 {cite("cohen1988")}. Meta-analytic estimates of this association are far larger (r ≈ 0.62) {cite("mcmullan2019","schenkel2021")}, so these targets are conservative. With the achieved sample of {n}, the study had {ss["power_r_0_20"]*100:.1f}% power to detect r = 0.20, and the smallest detectable correlation at 80% power was r = {ss["min_detectable_r"]:.2f}. For the multivariable model with 20 predictors, power to detect the unique contribution of a single predictor of small size (f² = 0.03) was {pow20*100:.1f}%. For descriptive proportions, {n} participants give a margin of error of ±{ss["moe_at_n"]:.1f}% at 95% confidence (worst case, p = 0.5), compared with ±5% for the {ss["cochran_n_for_5pct"]} participants of a Cochran-based calculation. The logistic model for clinically relevant health anxiety had {ss["HA_clin_events"]} events for 11 predictors (≈5 events per variable), which is acknowledged as a limitation of that model.')
H('2.4 Data collection instrument', 2)
P('A bilingual (Arabic–English) self-administered questionnaire was built in Google Forms. All items were required, except the free-text description of chronic disease and the name of the AI chatbot used, which were conditional. The questionnaire had five parts:')
numbered('**Consent and socio-demographic data:** confirmation of age ≥18 and consent; age group; gender; highest education; nature of study or work; diagnosed chronic medical condition (and its type).')
numbered('**Digital health information seeking:** tools used (search engines, AI chatbots, dedicated medical platforms, social media; multiple answers); main AI chatbot; motivation for searching (multiple answers); and trust in medical information generated by online tools/AI (5-point scale).')
numbered(f'**Health anxiety:** the 14-item Short Health Anxiety Inventory (SHAI-14) {cite("salk2002")}. Each item offers four statements scored 0–3, giving a total of 0–42; higher scores indicate greater health anxiety. The SHAI has good reliability and validity in clinical and non-clinical samples {cite("abramowitz2007","alberts2013")}, and an Arabic version has been validated in Saudi Arabia {cite("alshayea2020")}.')
numbered(f'**Cyberchondria:** the CSS-12 {cite("mcelroy2019")}, with items rated on a 5-point frequency scale (1 = never to 5 = always), giving a total of 12–60 and four subscales (3–15 each): Excessiveness (items 1–3), Compulsion (4–6), Distress (7–9) and Reassurance seeking (10–12).')
numbered('**Pre-consultation behaviour and relationship with physicians:** behavioural responses after searching (multiple answers), whether search or AI results are disclosed to the physician, the perceived effect of digital information seeking on the relationship with and trust in physicians, and the language mainly relied upon while answering.')
P(f'Each item was presented in English and Arabic. The research team produced the Arabic wording; its equivalence was examined empirically by comparing the psychometric behaviour of the scales between participants who relied mainly on Arabic and those who relied mainly on English (Section 2.6). Response options were displayed in the form in an order that occasionally differed from the original scales; items were therefore scored according to the meaning of the chosen statement, following the original scoring key, rather than its position on the form.')
H('2.5 Variables and operational definitions', 2)
for b in [f'**Clinically relevant health anxiety:** SHAI-14 ≥ 18, the optimal cut-off for distinguishing patients with health anxiety from healthy controls in a recent diagnostic accuracy study {cite("osterman2022")}. Prevalence at alternative cut-offs (15 and 22) was also reported.',
          '**High cyberchondria:** CSS-12 total in the upper tertile of the sample (> 33), since the CSS-12 has no validated clinical cut-off.',
          '**AI-chatbot user:** ticked "AI chatbots" as a tool or named an AI chatbot in the follow-up question. Thirty-two respondents named a chatbot without ticking the tool option; they were classified as users (the follow-up question was optional and specifically addressed to chatbot users).',
          '**Health-related background:** medical students, students in other health fields and healthcare workers.',
          '**Occupation:** free-text "Other" answers were recoded into the closest category (e.g. secondary-school students and school leavers as non-health students; housewives, engineers and business owners as non-health employed/other). Age groups 40–49, 50–59 and 60+ were merged (≥40) for inferential analysis because of small cell sizes.']:
    bullet(b)
H('2.6 Statistical analysis', 2)
P('Analyses were carried out in Python 3 (pandas, SciPy, statsmodels, pingouin, factor_analyzer and semopy). The analysis code is fully scripted and reproducible. All tests were two-sided with α = 0.05.')
for b in [
 f'**Descriptive statistics:** frequencies and percentages for categorical variables; mean ± SD and median (IQR) for scores. Normality was assessed with the Shapiro–Wilk test, skewness and kurtosis.',
 f'**Psychometrics:** internal consistency by Cronbach\'s α (with 95% CI) and McDonald\'s ω {cite("hayescoutts2020")}; corrected item–rest correlations; sampling adequacy (KMO, Bartlett\'s test); and confirmatory factor analysis using diagonally weighted least squares (appropriate for ordinal items) with maximum-likelihood estimates reported alongside as a stricter check. Fit was judged against conventional criteria (CFI/TLI ≥ 0.95, RMSEA ≤ 0.06–0.08) {cite("hu1999")}. Equivalence of the Arabic and English responders was tested by comparing α values with Feldt\'s test {cite("feldt1987")}, factor-loading similarity with Tucker\'s congruence coefficient {cite("lorenzo2006")}, and adjusted mean differences.',
 f'**Correlation analysis:** Pearson and Spearman correlations (bootstrapped 95% CIs) and partial correlations adjusted for covariates. Differences between dependent correlations (health anxiety with each CSS subscale) were tested with Williams\' t {cite("steiger1980")}. Item-level correlations were computed for every SHAI item with the CSS total and every CSS item with the SHAI total.',
 f'**Group comparisons for every variable:** Mann–Whitney U (rank-biserial r as effect size) or Kruskal–Wallis H (ε² as effect size) for SHAI, CSS and each CSS subscale, with Welch\'s t-test/ANOVA as parametric confirmation and pairwise Mann–Whitney post-hoc comparisons with Holm adjustment. Associations of all categorical variables with clinically relevant health anxiety and high cyberchondria were tested with Pearson χ² or Fisher\'s exact test (Monte-Carlo exact p when expected counts were small), with Cramér\'s V. Pairwise associations among all categorical variables were also screened.',
 f'**Multiple testing:** because many associations were tested, p-values within each family of tests were adjusted with the Benjamini–Hochberg false discovery rate procedure (q-values) {cite("bh1995")}. Unadjusted p-values are reported, and associations that do not survive correction are labelled as exploratory.',
 f'**Multivariable models:** hierarchical OLS regression of CSS on socio-demographics (block 1), digital-use variables (block 2) and SHAI (block 3), with HC3 robust standard errors, standardised coefficients and VIF checks; the reverse model (SHAI on CSS); models for each CSS subscale; quantile regression (25th–90th percentiles) and a quadratic term to examine non-linearity. Logistic regression identified predictors of clinically relevant health anxiety (AORs, AUC) and of each pre-consultation behaviour; ordinal logistic regression was used for the ordered disclosure and relationship outcomes.',
 f'**Mediation and moderation:** mediation models with bootstrap confidence intervals (5000 resamples) tested whether excessive searching mediated the link between health anxiety and post-search distress, and related pathways, including a reverse-direction model and a serial model {cite("hayes2022")}. Moderation of the SHAI–CSS association by trust, AI use, health background, gender, chronic disease, language and age was tested with interaction terms and simple slopes.',
 '**Sensitivity analyses:** exclusion of annotated responses; exclusion of straight-line responders (identical answers to all items of a scale); alternative SHAI cut-offs; bootstrap of the main regression coefficient; and comparison of early versus late responders as a proxy for non-response bias.']:
    bullet(b)
H('2.7 Ethical considerations', 2)
P('Participation was voluntary and anonymous. The first page of the form explained the purpose of the study, and electronic informed consent (with confirmation of age ≥18) was required before any question was shown; participants could withdraw at any time by closing the form. No names or contact details were collected, and data were stored securely and used only for research purposes. [Ethical approval: committee name, approval number and date to be inserted.]')

# ------------------------------------------------------------------ RESULTS
H('3. Results')
H('3.1 Characteristics of participants', 2)
occ = d.occupation.value_counts()
P(f'A total of {n} adults completed the questionnaire; all gave consent and no item had missing data. The sample was young, with {npct(age20)} aged 20–29 and {npct(int((d.age_group=="18-19").sum()))} aged 18–19; {npct(fem)} were female. Most were currently students ({npct(int((d.education=="Currently a student").sum()))}); medical students were the largest occupational group ({npct(medst)}), and {npct(hb)} had a health-related background. A diagnosed chronic condition was reported by {npct(cnt("chronic"))}, most commonly migraine, hypertension, asthma, anaemia, psoriasis and diabetes. Most participants relied mainly on Arabic when answering ({npct(int((d.language=="Arabic").sum()))}) (Table 1).')
T1 = X['T1_demographics']
rows = []
for var in T1.Variable.unique():
    rows.append([var, '', ''])
    for _, r in T1[T1.Variable == var].iterrows():
        cat = {'1': 'Yes', '0': 'No'}.get(str(r.Category), str(r.Category))
        rows.append(['    ' + cat, int(r.n), f'{r["%"]:.1f}'])
caption('Table', f'Socio-demographic characteristics of the participants (n = {n})')
table(['Characteristic', 'n', '%'], rows, widths=[10.5, 2, 2], bold_rows=[i for i, r in enumerate(rows) if r[1] == ''])

H('3.2 Patterns of online health information seeking', 2)
P(f'AI chatbots were the most frequently used tool for looking up symptoms ({npct(ai_n)}), followed by search engines ({npct(se_n)}), dedicated medical platforms ({npct(cnt("tool_medplat"))}) and social media ({npct(cnt("tool_social"))}); {npct(int((d.n_tools>1).sum()))} used more than one tool. ChatGPT was the dominant chatbot ({npct(int((d.ai_model=="ChatGPT").sum()))}), followed by Gemini ({npct(int((d.ai_model=="Gemini").sum()))}) and Claude ({npct(int((d.ai_model=="Claude").sum()))}). Notably, {npct(int((d.tool_profile=="AI without search engine").sum()))} used AI chatbots without using search engines at all. Convenience was the leading motive ({pct(cnt("mot_convenience"))}), followed by deciding whether a doctor visit was needed ({pct(cnt("mot_triage"))}) and seeking instant reassurance about an unexplained symptom ({pct(cnt("mot_reassurance"))}). Trust in medical information generated by online tools or AI was mostly moderate ({npct(int((d.trust==3).sum()))}); only {npct(int((d.trust>=4).sum()))} trusted it very much or completely (Table 2, Figure 1).')
T2 = X['T2_digital_behaviour']
rows = []
for var in T2.Variable.unique():
    rows.append([var, '', '']);
    for _, r in T2[T2.Variable == var].iterrows():
        cat = str(r.Category)
        if var.startswith('Trust'): cat = {'1': '1 – Do not trust at all', '2': '2 – Slightly', '3': '3 – Moderately', '4': '4 – Very much', '5': '5 – Completely'}[cat]
        rows.append(['    ' + cat, int(r.n), f'{r["%"]:.1f}'])
caption('Table', 'Online health information seeking, pre-consultation behaviour and relationship with physicians (n = 300)')
table(['Variable', 'n', '%'], rows, widths=[10.5, 2, 2], bold_rows=[i for i, r in enumerate(rows) if r[1] == ''],
      note='Tools, motivations and behaviours allowed multiple answers, so percentages do not sum to 100. AI users include 32 respondents who named a chatbot without ticking the AI option.')
figure(f'{FIG}/fig2_tools_motives.png', 'Tools used and motivations for searching for health information online (n = 300).')

H('3.3 Pre-consultation behaviour and the doctor–patient relationship', 2)
P(f'After searching, the most frequent response was to delay or skip professional care because the online information was reassuring ({npct(cnt("beh_delay_reassured"))}), followed by ignoring the symptom ({npct(cnt("beh_ignore"))}). Behaviours suggesting increased use of services were also common: consulting multiple doctors to double-check findings ({npct(cnt("beh_multidoc"))}) and requesting an urgent appointment or visiting the emergency room ({npct(cnt("beh_urgent"))}). Self-medication without consulting a doctor was reported by {npct(cnt("beh_selfmed"))}, and {npct(cnt("beh_delay_terrified"))} delayed care because the information had frightened or overwhelmed them. Only {npct(int((d.share=="Always").sum()))} always disclosed their online or AI findings to their physician, whereas {npct(cnt("share_never"))} never did so for fear of being judged or dismissed. Although most reported no negative effect on their relationship with physicians, {npct(int((d.relationship=="Somewhat complicated").sum()))} said it had become somewhat complicated and {npct(int((d.relationship=="Significantly reduced trust").sum()))} that it had significantly reduced their trust in their doctor (Figure 2).')
figure(f'{FIG}/fig3_behaviours.png', 'Behavioural responses after searching, disclosure of searches to physicians and perceived effect on the doctor–patient relationship (n = 300).', width=16)

H('3.4 Levels of health anxiety and cyberchondria, and psychometric properties', 2)
T3 = X['T3_scale_desc']
P(f'The mean SHAI-14 score was {msd(shai)} (median {shai.median():.0f}, IQR {shai.quantile(.25):.0f}–{shai.quantile(.75):.0f}). Using the cut-off of 18, {npct(cnt("HA_clin"))} (95% CI {ci(*prev["HA_clin_CI"], k=1)}%) had clinically relevant health anxiety; prevalence was {sens["prev_SHAI_ge_15"]:.1f}% and {sens["prev_SHAI_ge_22"]:.1f}% at cut-offs of 15 and 22, respectively. The mean CSS-12 score was {msd(css)} (median {css.median():.0f}). Excessiveness was the highest-scoring dimension ({msd(d.CSS_excess)}) and compulsion the lowest ({msd(d.CSS_compul)}). At item level, {X["T27_item_endorsement"].iloc[0].pct_often_always:.1f}% often or always searched online when they noticed an unexplained bodily sensation, whereas fewer than 7% reported that searching often interfered with work, leisure or social activities. Scores deviated modestly from normality (Shapiro–Wilk p < 0.01; skewness {T3.loc["SHAI-14 total (0-42)","skew"]:.2f} for SHAI and {T3.loc["CSS-12 total (12-60)","skew"]:.2f} for CSS), so non-parametric tests were used for bivariate comparisons and robust standard errors for regression.')
P(f'Both scales showed good reliability (Table 3): α = {rel.loc["SHAI-14","alpha"]:.2f} (ω = {rel.loc["SHAI-14","omega"]:.2f}) for SHAI-14 and α = {rel.loc["CSS-12","alpha"]:.2f} (ω = {rel.loc["CSS-12","omega"]:.2f}) for CSS-12, with subscale α between {min(rel.loc[k,"alpha"] for k in ["CSS Excessiveness","CSS Compulsion","CSS Distress","CSS Reassurance"]):.2f} and {max(rel.loc[k,"alpha"] for k in ["CSS Excessiveness","CSS Compulsion","CSS Distress","CSS Reassurance"]):.2f}. Sampling adequacy was good (KMO = {R["kmo"]["SHAI"]["KMO"]:.2f} and {R["kmo"]["CSS"]["KMO"]:.2f}; Bartlett p < 0.001). The four-factor structure of the CSS-12 fitted the data well (DWLS: CFI = {cfa.iloc[2].CFI:.3f}, RMSEA = {cfa.iloc[2].RMSEA:.3f}; ML: CFI = {cfa.iloc[3].CFI:.3f}, RMSEA = {cfa.iloc[3].RMSEA:.3f}) and clearly better than a single factor (ML CFI = {cfa.iloc[1].CFI:.3f}); standardised loadings ranged from {R["css4_std_loadings_range"][0]:.2f} to {R["css4_std_loadings_range"][1]:.2f}. For the SHAI-14, a one-factor model fitted well under DWLS (CFI = {cfa.iloc[4].CFI:.3f}, RMSEA = {cfa.iloc[4].RMSEA:.3f}) but only moderately under ML (CFI = {cfa.iloc[5].CFI:.3f}, RMSEA = {cfa.iloc[5].RMSEA:.3f}), and a two-factor model gave little improvement. Reliability was virtually identical in Arabic and English responders (SHAI α {R["feldt"]["SHAI"]["alpha_ar"]:.2f} vs {R["feldt"]["SHAI"]["alpha_en"]:.2f}, Feldt p = {R["feldt"]["SHAI"]["p"]:.2f}; CSS α {R["feldt"]["CSS"]["alpha_ar"]:.2f} vs {R["feldt"]["CSS"]["alpha_en"]:.2f}, p = {R["feldt"]["CSS"]["p"]:.2f}), factor loadings were highly congruent (Tucker φ = {R["tucker_congruence"]["SHAI"]:.2f} and {R["tucker_congruence"]["CSS"]:.2f}), and mean scores did not differ by language after adjustment (Table 3).')
rows = []
for lab, col, rk in [('SHAI-14 total (0–42)', 'SHAI', 'SHAI-14'), ('CSS-12 total (12–60)', 'CSS', 'CSS-12'), ('  Excessiveness (3–15)', 'CSS_excess', 'CSS Excessiveness'),
                     ('  Compulsion (3–15)', 'CSS_compul', 'CSS Compulsion'), ('  Distress (3–15)', 'CSS_distress', 'CSS Distress'), ('  Reassurance (3–15)', 'CSS_reassur', 'CSS Reassurance')]:
    s = d[col]; rr = rel.loc[rk]
    rows.append([lab, f'{s.mean():.2f} ± {s.std():.2f}', f'{s.median():.0f} ({s.quantile(.25):.0f}–{s.quantile(.75):.0f})', f'{s.min():.0f}–{s.max():.0f}', f'{rr.alpha:.2f} ({rr.alpha_CI})', f'{rr.omega:.2f}'])
lr = lang.set_index('Scale')
caption('Table', 'Descriptive statistics, reliability and language equivalence of the SHAI-14 and CSS-12 (n = 300)')
table(['Scale', 'Mean ± SD', 'Median (IQR)', 'Range', 'α (95% CI)', 'ω'], rows, widths=[4.6, 2.6, 2.6, 1.6, 2.7, 1.1],
      note=f'Language equivalence: SHAI mean {lr.loc["SHAI","Arabic"]} (Arabic) vs {lr.loc["SHAI","English"]} (English), adjusted p = {lr.loc["SHAI","adj_p"]:.2f}; CSS mean {lr.loc["CSS","Arabic"]} vs {lr.loc["CSS","English"]}, adjusted p = {lr.loc["CSS","adj_p"]:.2f}. CFA fit (CSS-12, four factors): DWLS CFI {cfa.iloc[2].CFI:.3f}, TLI {cfa.iloc[2].TLI:.3f}, RMSEA {cfa.iloc[2].RMSEA:.3f}; ML CFI {cfa.iloc[3].CFI:.3f}, TLI {cfa.iloc[3].TLI:.3f}, RMSEA {cfa.iloc[3].RMSEA:.3f}. SHAI-14 one factor: DWLS CFI {cfa.iloc[4].CFI:.3f}, RMSEA {cfa.iloc[4].RMSEA:.3f}; ML CFI {cfa.iloc[5].CFI:.3f}, RMSEA {cfa.iloc[5].RMSEA:.3f}.')

H('3.5 Association between health anxiety and cyberchondria', 2)
P(f'Health anxiety was positively and moderately correlated with cyberchondria (Pearson r = {sc["pearson_r"]:.2f}, 95% CI {ci(*sc["pearson_CI"])}; Spearman ρ = {sc["spearman"]:.2f}, bootstrap 95% CI {ci(*sc["spearman_boot_CI"])}; {pv(sc["pearson_p"])}), with health anxiety accounting for {sc["r2"]*100:.0f}% of the variance in cyberchondria (Figure 3). The association was essentially unchanged after adjustment for age, gender, health background, chronic disease, trust in online information and AI use (partial r = {part["r"]:.2f}, 95% CI {ci(*part["CI"])}). Health anxiety was correlated with all four cyberchondria dimensions, but not equally: correlations with excessiveness (r = {r_sub["Excessiveness"]["pearson"]:.2f}) and distress (r = {r_sub["Distress"]["pearson"]:.2f}) were similar to each other and significantly stronger than those with compulsion (r = {r_sub["Compulsion"]["pearson"]:.2f}) and reassurance seeking (r = {r_sub["Reassurance"]["pearson"]:.2f}) (Williams\' t, all p < 0.001) (Table 4, Figure 4).')
P(f'At item level, the CSS items most strongly linked to health anxiety were panicking when reading that a symptom could signal a rare or serious condition (ρ = {X["T8_item_level_corr"].iloc[6].rho:.2f}), searching whenever an unexplained sensation is noticed (ρ = {X["T8_item_level_corr"].iloc[0].rho:.2f}) and feeling more anxious after searching (ρ = {X["T8_item_level_corr"].iloc[8].rho:.2f}). Among SHAI items, fear of serious illness (ρ = {X["T8_item_level_corr"].iloc[16].rho:.2f}) and interpretation of bodily sensations (ρ = {X["T8_item_level_corr"].iloc[21].rho:.2f}) correlated most with cyberchondria, whereas simple awareness of bodily sensations and of aches and pains did not (ρ = {X["T8_item_level_corr"].iloc[14].rho:.2f} and {X["T8_item_level_corr"].iloc[13].rho:.2f}, q > 0.10). The more tools and motives participants reported, the higher both scores were, whereas trust in online/AI information was associated only with excessiveness (ρ = {X["T7_corr_matrix_rho"].loc["Trust in online/AI","Excessiveness"]:.2f}). A full item-by-item analysis of all 168 SHAI × CSS item pairs (Appendix, Figure 9) showed that {ib["sig_fdr"]} pairs remained significant after FDR correction and {ib["sig_holm"]} after Holm–Bonferroni correction; the strongest were fear of serious illness with panic on reading about a rare or serious condition (ρ = {ib["top"][0][2]:.2f}), interpreting bodily sensations with searching whenever a sensation is noticed (ρ = {ib["top"][1][2]:.2f}), and difficulty distracting oneself from health thoughts with the same panic item (ρ = {ib["top"][2][2]:.2f}). Simple awareness of aches, pains and bodily sensations (SHAI items 2 and 3) was largely unrelated to any searching behaviour.')
rows = []
for k, col in [('CSS-12 total', None), ('Excessiveness', 'Excessiveness'), ('Compulsion', 'Compulsion'), ('Distress', 'Distress'), ('Reassurance', 'Reassurance')]:
    if col is None:
        rows.append([k, f'{sc["pearson_r"]:.2f} ({ci(*sc["pearson_CI"])})', f'{sc["spearman"]:.2f}', pnum(sc["pearson_p"]), f'{b_shai.beta_std:.2f}'])
    else:
        rr = r_sub[col]; sq = subq.set_index('Outcome').loc[col]
        rows.append([k, f'{rr["pearson"]:.2f} ({ci(*rr["pearson_CI"])})', f'{rr["spearman"]:.2f}', pnum(rr["p"]), f'{sq.beta:.2f}'])
caption('Table', 'Correlation between health anxiety (SHAI-14) and cyberchondria (CSS-12) and its dimensions (n = 300)')
table(['Cyberchondria measure', 'Pearson r (95% CI)', 'Spearman ρ', 'p', 'Adjusted β'], rows, widths=[4.5, 4, 2.4, 2, 2.4],
      note='Adjusted β: standardised coefficient of SHAI from robust OLS models adjusted for age, gender, occupation, chronic disease, language, tools, trust and motives. Excessiveness and Distress correlations were each significantly stronger than Compulsion and Reassurance (Williams\' t, p < 0.001); they did not differ from each other (p = ' + f'{cmpc.iloc[1].p:.2f}).')
figure(f'{FIG}/fig1_scatter_shai_css.png', 'Scatter plot of health anxiety (SHAI-14) against cyberchondria (CSS-12) with the least-squares line and 95% confidence band (n = 300; points slightly jittered).', width=13)
figure(f'{FIG}/fig4_corr_heatmap.png', 'Spearman correlation matrix of health anxiety, cyberchondria and its dimensions, trust in online/AI information, number of tools, number of motives and age group (FDR-adjusted significance).', width=14)

H('3.6 Factors associated with health anxiety and cyberchondria (bivariate analysis)', 2)
gA = gp('Age group', 'SHAI'); gE = gp('Education', 'SHAI'); gO = gp('Study/work', 'SHAI'); gC = gp('Chronic condition', 'SHAI')
gAIc = gp('Uses AI chatbots', 'CSS'); gAIs = gp('Uses AI chatbots', 'SHAI'); gSMs = gp('Uses social media', 'SHAI'); gSMc = gp('Uses social media', 'CSS')
gTPc = gp('Tool profile', 'CSS'); gMR = gp('Motive: reassurance', 'SHAI'); gGs = gp('Gender', 'SHAI'); gGc = gp('Gender', 'CSS')
P(f'Every variable was tested against SHAI, CSS and each CSS dimension (162 tests; full results in the supplementary workbook). Table 5 summarises the main comparisons. Health anxiety was higher in younger participants (18–19 years: {msd(d.loc[d.age4=="18-19","SHAI"])} vs 30–39 years: {msd(d.loc[d.age4=="30-39","SHAI"])}; {pv(gA.p)}, q = {gA.q_FDR:.3f}), in current students ({pv(gE.p)}), in medical students compared with non-health employees ({pv(gO.p)}), in those with a chronic condition ({pv(gC.p)}, q = {gC.q_FDR:.3f}), in social-media users ({pv(gSMs.p)}, q = {gSMs.q_FDR:.3f}) and, most clearly, in those whose motive was to obtain instant reassurance ({msd(d.loc[d.mot_reassurance==1,"SHAI"])} vs {msd(d.loc[d.mot_reassurance==0,"SHAI"])}; {pv(gMR.p)}, q = {gMR.q_FDR:.3f}). Gender was not associated with either health anxiety ({pv(gGs.p)}) or cyberchondria ({pv(gGc.p)}).')
P(f'Cyberchondria was higher among AI-chatbot users than non-users ({msd(d.loc[d.ai_user==1,"CSS"])} vs {msd(d.loc[d.ai_user==0,"CSS"])}; {pv(gAIc.p)}, rank-biserial r = {gAIc.ES:.2f}, q = {gAIc.q_FDR:.3f}) and differed by tool profile, rising from non-AI users to those who used AI alone and was highest in those who combined AI with search engines ({pv(gTPc.p)}; Figure 5). Of the CSS dimensions, AI use was linked with excessiveness and reassurance seeking but not with compulsion. Among chatbot users, the specific chatbot (ChatGPT, Gemini or Claude) made no difference to cyberchondria ({pv(kw_bot)}). Participants who delayed care because online information terrified them had higher scores on every measure, particularly distress ({pv(gp("Delay care (terrified)","Distress").p)}), whereas those who decided to ignore the symptom had lower cyberchondria ({pv(gp("Ignore symptom","CSS").p)}). Participants who reported urgent visits had higher reassurance-seeking scores ({pv(gp("Urgent visit/ER","Reassurance").p)}).')
# Table 5: selected comparisons
sel = [('Age group', 'age4', ['18-19', '20-29', '30-39', '≥40']), ('Gender', 'gender', ['Female', 'Male']),
       ('Study/work', 'occupation', ['Medical student', 'Other health-science student', 'Healthcare worker', 'Non-health student', 'Non-health employee/other']),
       ('Chronic condition', 'chronic', [1, 0]), ('Uses AI chatbots', 'ai_user', [1, 0]), ('Uses search engines', 'tool_search', [1, 0]),
       ('Uses social media', 'tool_social', [1, 0]), ('Tool profile', 'tool_profile', ['AI + search engine', 'AI without search engine', 'No AI (search/other only)']),
       ('Motive: reassurance', 'mot_reassurance', [1, 0]), ('Motive: self-diagnosis', 'mot_selfdx', [1, 0]),
       ('Trust in online/AI (1-5)', 'trust', [1, 2, 3, 4, 5]), ('Language', 'language', ['Arabic', 'English'])]
rows = []; bolds = []
for lab, col, levels in sel:
    g1 = gp(lab, 'SHAI'); g2 = gp(lab, 'CSS')
    bolds.append(len(rows))
    rows.append([lab.replace(' (1-5)', ''), '', '', f'{pnum(g1.p)} ({g1.q_FDR:.3f})', '', f'{pnum(g2.p)} ({g2.q_FDR:.3f})'])
    for lv in levels:
        s = d[d[col] == lv]
        name = {1: 'Yes', 0: 'No'}.get(lv, lv) if col not in ('trust',) else str(lv)
        rows.append(['    ' + str(name), len(s), f'{s.SHAI.mean():.1f} ± {s.SHAI.std():.1f}', '', f'{s.CSS.mean():.1f} ± {s.CSS.std():.1f}', ''])
caption('Table', 'Health anxiety (SHAI-14) and cyberchondria (CSS-12) scores by participant characteristics (n = 300)')
table(['Variable', 'n', 'SHAI mean ± SD', 'p (q)', 'CSS mean ± SD', 'p (q)'], rows, widths=[5.2, 1.2, 2.8, 2.4, 2.8, 2.4], size=10, bold_rows=bolds,
      note='p: Mann–Whitney U (two groups) or Kruskal–Wallis H (three or more groups); q: Benjamini–Hochberg false-discovery-rate adjusted p across all 162 group comparisons. Effect sizes, medians (IQR), Welch tests and post-hoc comparisons are provided in the supplementary workbook.')
figure(f'{FIG}/fig6_css_by_tool.png', 'Cyberchondria (CSS-12) by online tool profile.', width=12)

P(f'In categorical analysis, clinically relevant health anxiety was more common among social-media users ({chi("Clinically","Uses social media").Distribution.split("; ")[1].split(": ")[1]} vs {chi("Clinically","Uses social media").Distribution.split("; ")[0].split(": ")[1]}; q = {chi("Clinically","Uses social media").q_FDR:.3f}) and those seeking instant reassurance ({chi("Clinically","Motive: reassurance").Distribution.split("; ")[1].split(": ")[1]} vs {chi("Clinically","Motive: reassurance").Distribution.split("; ")[0].split(": ")[1]}; q = {chi("Clinically","Motive: reassurance").q_FDR:.3f}). Participants with clinically relevant health anxiety had 3.7 times the odds of high cyberchondria (52.6% vs 23.0%; OR = {R["HA_CSShigh_OR"]["OR"]:.2f}, 95% CI {ci(*R["HA_CSShigh_OR"]["CI"])}). High cyberchondria was associated with disclosure pattern ({pv(chi("High","Disclose to physician").p)}, q = {chi("High","Disclose to physician").q_FDR:.3f}): it was most frequent among those who always disclosed their searches ({chi("High","Disclose to physician").Distribution.split("; ")[0].split(": ")[1]}) and least frequent among those who never did ({chi("High","Disclose to physician").Distribution.split("; ")[1].split(": ")[1]}) (Table 6).')
rows = []
for out_lab, out in [('Clinically relevant HA (SHAI ≥ 18)', 'Clinically'), ('High cyberchondria (CSS > 33)', 'High')]:
    sub = CH[CH.Outcome.str.startswith(out) & (CH.p < 0.05)].sort_values('p')
    rows.append([out_lab, '', '', '', ''])
    for _, r in sub.iterrows():
        dist = r.Distribution.replace('Clinically relevant HA', '').replace('(fear of judgement)', '')
        dist = re.sub(r'\b0:', 'No:', dist); dist = re.sub(r'\b1:', 'Yes:', dist)
        rows.append(['    ' + r.Variable, dist, f'{r.CramersV:.2f}', pnum(r.p), f'{r.q_FDR:.3f}'])
caption('Table', 'Variables significantly associated (p < 0.05) with clinically relevant health anxiety and with high cyberchondria')
table(['Variable', 'n/N (%) with outcome, by group', 'Cramér\'s V', 'p', 'q (FDR)'], rows, widths=[4.6, 7.4, 1.6, 1.4, 1.6], size=9.5,
      bold_rows=[i for i, r in enumerate(rows) if r[1] == ''], note='Pearson χ² or Fisher\'s exact test (Monte-Carlo exact p where expected counts were < 5). Associations with q ≥ 0.05 are exploratory. Non-significant associations are listed in the supplementary workbook.')

H('3.7 Independent predictors (multivariable analysis)', 2)
P(f'In the hierarchical regression (Table 7), socio-demographic variables alone explained little of the variance in cyberchondria (R² = {hier.R2.iloc[0]:.3f}). Adding digital-use variables raised R² to {hier.R2.iloc[1]:.3f} (ΔR² = {hier.dR2.iloc[1]:.3f}, {pv(hier.p_change.iloc[1])}), and adding health anxiety raised it to {hier.R2.iloc[2]:.3f} (ΔR² = {hier.dR2.iloc[2]:.3f}, {pv(hier.p_change.iloc[2])}; adjusted R² = {hier.adjR2.iloc[2]:.3f}). In the final model, health anxiety was by far the strongest predictor (B = {b_shai.B:.2f}, 95% CI {ci(b_shai.CI_low, b_shai.CI_high)}, β = {b_shai.beta_std:.2f}, {pv(b_shai.p)}): each additional SHAI point was associated with a {b_shai.B:.2f}-point higher CSS score. AI-chatbot use was independently associated with a {b_ai.B:.1f}-point higher CSS score (95% CI {ci(b_ai.CI_low, b_ai.CI_high)}, β = {b_ai.beta_std:.2f}, {pv(b_ai.p)}), and age 30–39 years with higher cyberchondria than age 20–29 ({pv(ols.loc[ols.index.str.contains("30–39")].p.iloc[0])}). This AI finding was sensitive to how AI use was defined: when only respondents who ticked "AI chatbots" were counted as users (n = 234), so that the 32 who named a chatbot without ticking it fell into the comparison group, the association was no longer significant (B = {ai_alt_B:.2f}, {pv(ai_alt_p)}). The 32 reclassified respondents had the highest cyberchondria of any group (mean {d.loc[d.ai_model_inconsistent==1,"CSS"].mean():.1f}), and the true non-user group was small (n = 34). Occupation, health background, gender, chronic disease, language, trust and motives were not independently associated with cyberchondria. Multicollinearity was low (maximum VIF = {X["T14b_VIF"].VIF.max():.1f}), and the bootstrap 95% CI for the SHAI coefficient ({ci(*sens["boot_B_SHAI_CI"])}) confirmed its stability. Quantile regression showed a consistent effect of health anxiety across the distribution of cyberchondria (B from {quant.B_SHAI.min():.2f} to {quant.B_SHAI.max():.2f}), and a quadratic term was not significant ({pv(R["quadratic"]["p"])}), indicating a linear association.')
P(f'In the reverse model, cyberchondria was the strongest predictor of health anxiety (β = {hiS.beta_std:.2f}, {pv(hiS.p)}; ΔR² = {R["ols_shai"]["dR2_CSS"]:.3f}), together with chronic disease (B = {hiChron.B:.2f}, {pv(hiChron.p)}) and searching for instant reassurance (B = {hiMot.B:.2f}, {pv(hiMot.p)}); after adjustment, medical students no longer differed from other groups ({pv(olsS.loc[olsS.index.str.contains("Medical student")].p.iloc[0])}).')
rows = []
lab = {'Intercept': 'Constant', "C(age_c, Treatment('20–29'))[T.18–19]": 'Age 18–19 (ref. 20–29)', "C(age_c, Treatment('20–29'))[T.30–39]": 'Age 30–39', "C(age_c, Treatment('20–29'))[T.≥40]": 'Age ≥40',
       "C(occ_c, Treatment('Non-health employed/other'))[T.Healthcare worker]": 'Healthcare worker (ref. non-health employed)', "C(occ_c, Treatment('Non-health employed/other'))[T.Medical student]": 'Medical student',
       "C(occ_c, Treatment('Non-health employed/other'))[T.Non-health student]": 'Non-health student', "C(occ_c, Treatment('Non-health employed/other'))[T.Other health student]": 'Other health-science student',
       'female': 'Female', 'chronic': 'Chronic condition', 'lang_en': 'Answered mainly in English', 'ai_user': 'Uses AI chatbots', 'tool_search': 'Uses search engines',
       'tool_medplat': 'Uses medical platforms', 'tool_social': 'Uses social media', 'trust': 'Trust in online/AI (per point)', 'mot_reassurance': 'Motive: instant reassurance',
       'mot_selfdx': 'Motive: self-diagnosis', 'mot_cost': 'Motive: cost/logistics', 'mot_triage': 'Motive: need for a doctor?', 'SHAI': 'Health anxiety (SHAI, per point)'}
for idx, r in ols.iterrows():
    rows.append([lab[idx], f'{r.B:.2f}', ci(r.CI_low, r.CI_high), '' if np.isnan(r.beta_std) else f'{r.beta_std:.2f}', pnum(r.p)])
caption('Table', 'Hierarchical robust linear regression of cyberchondria (CSS-12 total) – final model (n = 300)')
table(['Predictor', 'B', '95% CI', 'β', 'p'], rows, widths=[7.6, 1.6, 3.2, 1.4, 1.6], size=10,
      bold_rows=[i for i, r in enumerate(rows) if r[4] != '' and r[4] != 'nan' and (r[4] == '<0.001' or float(r[4]) < 0.05) and r[0] != 'Constant'],
      note=f'Block 1 (socio-demographics) R² = {hier.R2.iloc[0]:.3f}; Block 2 (+ digital use) R² = {hier.R2.iloc[1]:.3f}, ΔR² = {hier.dR2.iloc[1]:.3f}, p = {hier.p_change.iloc[1]:.4f}; Block 3 (+ SHAI) R² = {hier.R2.iloc[2]:.3f}, ΔR² = {hier.dR2.iloc[2]:.3f}, p < 0.001; adjusted R² = {hier.adjR2.iloc[2]:.3f}; F = {R["ols_css"]["F"]:.2f}, p < 0.001. HC3 robust standard errors. β: standardised coefficient.')

P(f'In logistic regression (Table 8, Figure 6), clinically relevant health anxiety was independently associated with higher cyberchondria (AOR = {lg_css.AOR:.2f} per SD, 95% CI {ci(lg_css.CI_low, lg_css.CI_high)}), with searching for instant reassurance (AOR = {lg_mot.AOR:.2f}, 95% CI {ci(lg_mot.CI_low, lg_mot.CI_high)}) and with younger age (AOR = {lg_age.AOR:.2f} per age-group step, 95% CI {ci(lg_age.CI_low, lg_age.CI_high)}). The model discriminated well (AUC = {R["logit_HA"]["AUC"]:.2f}; McFadden pseudo-R² = {R["logit_HA"]["pseudoR2"]:.2f}).')
rows = []
lab2 = {'age_ord': 'Age group (per step)', 'female': 'Female', 'health_background': 'Health-related background', 'chronic': 'Chronic condition', 'lang_en': 'Answered mainly in English',
        'ai_user': 'Uses AI chatbots', 'tool_search': 'Uses search engines', 'trust': 'Trust in online/AI (per point)', 'mot_reassurance': 'Motive: instant reassurance',
        'mot_selfdx': 'Motive: self-diagnosis', 'zCSS': 'Cyberchondria (per 1 SD)'}
for idx, r in lg.drop(index='Intercept').iterrows():
    rows.append([lab2[idx], f'{r.AOR:.2f}', ci(r.CI_low, r.CI_high), pnum(r.p)])
caption('Table', 'Multivariable logistic regression of clinically relevant health anxiety (SHAI ≥ 18; 57 events, n = 300)')
table(['Predictor', 'AOR', '95% CI', 'p'], rows, widths=[8, 2, 3.5, 2], bold_rows=[i for i, r in enumerate(rows) if r[3] == '<0.001' or float(r[3]) < 0.05],
      note=f'AUC = {R["logit_HA"]["AUC"]:.2f}; McFadden pseudo-R² = {R["logit_HA"]["pseudoR2"]:.2f}; likelihood-ratio p < 0.001; robust (HC1) standard errors.')
figure(f'{FIG}/fig7_forest_HA.png', 'Adjusted odds ratios for clinically relevant health anxiety (SHAI ≥ 18). Orange: p < 0.05.', width=14)

H('3.8 Health anxiety, cyberchondria and pre-consultation behaviour', 2)
bU = bh('Urgent visit/ER', 'zCSS'); bI = bh('Ignore symptom', 'zCSS'); bT = bh('Delay care (terrified)', 'zCSS'); bNS = bh('Never disclose to physician', 'zSHAI'); bNC = bh('Never disclose to physician', 'zCSS')
bM = bh('Multiple doctors', 'trust'); bR = bh('Negative effect on doctor relationship', 'zCSS'); oR = ordl[(ordl.Outcome.str.startswith('Damage')) & (ordl.Predictor == 'zCSS')].iloc[0]
P(f'When health anxiety and cyberchondria were entered together, each with adjustment for trust, AI use and socio-demographics (Table 9), cyberchondria, rather than health anxiety, was associated with behaviours: each SD increase in CSS raised the odds of an urgent visit or ER attendance (AOR = {bU.AOR:.2f}, 95% CI {ci(bU.CI_low, bU.CI_high)}) and of delaying care because of fear (AOR = {bT.AOR:.2f}, 95% CI {ci(bT.CI_low, bT.CI_high)}), and lowered the odds of ignoring the symptom (AOR = {bI.AOR:.2f}, 95% CI {ci(bI.CI_low, bI.CI_high)}). Higher cyberchondria was also associated with a more strained relationship with physicians (ordinal OR = {oR.OR:.2f}, 95% CI {ci(oR.CI_low, oR.CI_high)}). Strikingly, health anxiety and cyberchondria pulled in opposite directions for disclosure: higher health anxiety increased the odds of never telling the physician about online searches (AOR = {bNS.AOR:.2f}, 95% CI {ci(bNS.CI_low, bNS.CI_high)}), whereas higher cyberchondria decreased them (AOR = {bNC.AOR:.2f}, 95% CI {ci(bNC.CI_low, bNC.CI_high)}, q = {bNC.q_FDR:.3f}). Greater trust in online/AI information was associated with lower odds of consulting multiple doctors (AOR = {bM.AOR:.2f}, 95% CI {ci(bM.CI_low, bM.CI_high)}). Neither score was associated with self-medication or with delaying care because of reassurance. Apart from the disclosure finding, these behavioural associations did not survive FDR correction across the 40 tests (q between 0.08 and 0.28) and should be regarded as exploratory.')
rows = []
for outc in ['Urgent visit/ER', 'Multiple doctors', 'Delay/skip care (reassured)', 'Delay care (terrified)', 'Self-medication', 'Ignore symptom', 'Never disclose to physician', 'Negative effect on doctor relationship']:
    a = bh(outc, 'zSHAI'); b = bh(outc, 'zCSS')
    rows.append([outc, int(a.events), f'{a.AOR:.2f} ({ci(a.CI_low, a.CI_high)})', pnum(a.p), f'{b.AOR:.2f} ({ci(b.CI_low, b.CI_high)})', pnum(b.p)])
caption('Table', 'Adjusted associations of health anxiety and cyberchondria with pre-consultation behaviours and the doctor–patient relationship')
table(['Outcome', 'Events', 'SHAI AOR (95% CI)', 'p', 'CSS AOR (95% CI)', 'p'], rows, widths=[5, 1.4, 3.4, 1.4, 3.4, 1.4], size=10,
      note='AOR per 1 SD of each score, from logistic models containing both scores plus trust, AI use, age, gender, health background, chronic disease and language. FDR-adjusted q-values are given in the supplementary workbook.')

H('3.9 Pathways: mediation and moderation', 2)
P(f'Mediation analyses (standardised, adjusted, 5000 bootstrap resamples) supported pathways consistent with a reinforcing cycle (Table 10, Figure 7). First, excessive searching partly mediated the association between health anxiety and post-search distress (indirect effect = {mA_ind.coef:.2f}, 95% CI {ci(mA_ind.CI_low, mA_ind.CI_high)}; {pmed[[k for k in pmed if k.startswith("A:")][0]]*100:.0f}% of the total effect), with the direct effect remaining significant (c′ = {mA_dir.coef:.2f}). Second, post-search distress almost fully mediated the association between health anxiety and reassurance-seeking consultations (indirect = {mC_ind.coef:.2f}, 95% CI {ci(mC_ind.CI_low, mC_ind.CI_high)}; direct c′ = {mC_dir.coef:.2f}, {pv(mC_dir.p)}). A serial model showed a significant chain from health anxiety to excessive searching to distress to reassurance seeking (indirect = {ser.iloc[2].coef:.2f}, 95% CI {ci(ser.iloc[2].CI_low, ser.iloc[2].CI_high)}). The reverse ordering, from excessive searching through distress to health anxiety, was also significant (indirect = {mD_ind.coef:.2f}, 95% CI {ci(mD_ind.CI_low, mD_ind.CI_high)}), which is what a bidirectional relationship would produce, but also what cross-sectional data cannot disentangle. Participants with clinically relevant health anxiety scored higher on all four dimensions, most markedly on excessiveness and distress (Figure 8).')
P(f'None of the tested moderators significantly altered the strength of the SHAI–CSS association after FDR correction (Table 10): trust in online/AI information (interaction p = {mod.iloc[0].p:.2f}), AI use (p = {mod.iloc[1].p:.2f}), health background (p = {mod.iloc[2].p:.2f}), gender (p = {mod.iloc[3].p:.2f}), chronic disease (p = {mod.iloc[4].p:.2f}) and language (p = {mod.iloc[5].p:.2f}); there was a non-significant trend towards a weaker association with increasing age (p = {mod.iloc[6].p:.3f}). Consistently, the correlation was moderate to strong in every subgroup with adequate size, ranging from r = {sg[sg.n>=40].r.min():.2f} to {sg[sg.n>=40].r.max():.2f}; it was somewhat stronger in medical students (r = {R["medstud_corr"]["r_med"]:.2f}) than in other participants (r = {R["medstud_corr"]["r_other"]:.2f}; z-test p = {R["medstud_corr"]["p"]:.3f}).')
rows = []
for lab_, pre in [('A. HA → excessive searching → post-search distress', 'A:'), ('B. HA → excessive searching → compulsion', 'B:'),
                  ('C. HA → post-search distress → reassurance seeking', 'C:'), ('D. Excessive searching → distress → HA (reverse)', 'D:')]:
    a = med[med.Model.str.startswith(pre)].reset_index(drop=True)
    g = lambda pth: a[a.Path == pth].iloc[0]
    rows.append([lab_, f'{a.iloc[0].coef:.2f}', f'{a.iloc[1].coef:.2f}', f'{g("Total").coef:.2f}', f'{g("Direct").coef:.2f}{"" if g("Direct").p < .05 else " (ns)"}',
                 f'{g("Indirect").coef:.2f} ({ci(g("Indirect").CI_low, g("Indirect").CI_high)})'])
rows.append(['E. Serial: HA → excess → distress → reassurance', '', '', '', f'{ser.iloc[3].coef:.2f} (ns)', f'{ser.iloc[2].coef:.2f} ({ci(ser.iloc[2].CI_low, ser.iloc[2].CI_high)})'])
caption('Table', 'Mediation models linking health anxiety (HA) and the dimensions of cyberchondria (standardised coefficients, n = 300)')
table(['Model', 'a', 'b', 'Total (c)', 'Direct (c′)', 'Indirect (95% boot CI)'], rows, widths=[6.2, 1.2, 1.2, 1.7, 1.9, 3.5], size=10,
      note='All models adjusted for age, gender, health background, chronic disease, trust and AI use; 5000 bootstrap resamples. All a, b and total paths p < 0.001. Moderation of the SHAI–CSS association (interaction terms): ' +
      '; '.join([f'{r.Moderator.lower()} p = {r.p:.2f}' for _, r in mod.iterrows()]) + f' (all q > 0.40).')
figure(f'{FIG}/fig8_mediation.png', 'Mediation models: (A) excessive searching as a pathway from health anxiety to post-search distress; (B) post-search distress as a pathway from health anxiety to reassurance seeking.', width=16)
figure(f'{FIG}/fig5_subscales_by_HA.png', 'Cyberchondria dimensions by health-anxiety status (mean ± 95% CI).', width=13)

H('3.10 Sensitivity analyses', 2)
P(f'Results were robust. Excluding the three annotated responses (r = {sens["excluding_flagged_r"]:.2f}) or the {sens["straightliners"]} straight-line responders (n = {sens["n_after"]}, r = {sens["excluding_straightliners_r"]:.2f}) left the main correlation unchanged. Early and late responders did not differ in cyberchondria (p = {sens["early_late_CSS_p"]:.2f}) or health anxiety (p = {sens["early_late_SHAI_p"]:.2f}). The regression residuals showed mild non-normality and heteroskedasticity (Shapiro–Wilk p = {R["ols_css"]["resid_shapiro_p"]:.3f}; Breusch–Pagan p = {R["ols_css"]["bp_p"]:.3f}), which is why robust standard errors and bootstrap intervals were used; conclusions did not change.')

# ------------------------------------------------------------------ DISCUSSION
H('4. Discussion')
H('4.1 Principal findings', 2)
P(f'This study is, to our knowledge, the first from Iraq to examine health anxiety and cyberchondria together, and among the first anywhere to place AI chatbots alongside search engines as a source of pre-consultation health information. Four findings stand out. First, AI chatbots have already overtaken search engines as the main tool used to look up symptoms in this sample, and about two in five participants used AI without search engines at all. Second, health anxiety and cyberchondria were moderately and robustly associated (r ≈ 0.47), independently of all other measured factors, and health anxiety was the dominant predictor of cyberchondria. Third, the association was concentrated in the excessiveness and distress dimensions, and mediation analyses traced a sequence in which health anxiety is linked with excessive searching, excessive searching with post-search distress, and distress with reassurance-seeking consultations. Fourth, these psychological processes were reflected in pre-consultation behaviour, from urgent visits to non-disclosure of searches to physicians.')
H('4.2 Health anxiety and cyberchondria', 2)
P(f'The observed correlation (r = {sc["pearson_r"]:.2f}) confirms H1 and H2 but is weaker than the pooled estimates of about 0.62–0.63 from meta-analyses {cite("mcmullan2019","schenkel2021")}. Several explanations are plausible. Our sample had relatively low health anxiety (mean SHAI {shai.mean():.1f}) and a narrow range in which correlations are attenuated; recruitment through online channels may have favoured confident, habitual Internet users for whom searching is routine rather than alarming; and most previous studies used the full 33-item CSS, whose distress and reassurance items overlap more with health anxiety. The pattern across dimensions is also informative. Distress and excessiveness, the dimensions closest to the cognitive model of health anxiety, were most strongly related to SHAI, while compulsion (interference with work and social life) was rare and only weakly related, consistent with the subscale range reported previously {cite("schenkel2021")}. At item level, it was the *interpretation* of sensations and the *fear* of serious illness, not mere awareness of bodily sensations, that tracked cyberchondria. This supports the view that cyberchondria is driven by catastrophic appraisal rather than by bodily vigilance itself {cite("salk1986","fergus2013")}.')
P(f'The prevalence of clinically relevant health anxiety ({prev["HA_clin"]:.0f}% at SHAI ≥ 18) should be interpreted cautiously. The cut-off was derived against healthy controls in Sweden {cite("osterman2022")}; it is not a diagnosis, and prevalence ranged from 7% to 33% across plausible cut-offs. Nevertheless, it indicates that a substantial minority of digitally active adults in Baghdad experience health worry at a level that would warrant clinical attention, consistent with high rates of health anxiety reported in Iraq during the pandemic {cite("karim2020")}.')
H('4.3 Do health anxiety and cyberchondria fuel each other?', 2)
P(f'The mediation results are consistent with the reinforcing cycle proposed by cognitive-behavioural models and by the longitudinal Dutch study {cite("salk1986","tepoel2016")}: people with higher health anxiety search more, more searching goes with more distress, and distress goes with more reassurance seeking from doctors. That the direct effect of health anxiety on reassurance seeking disappeared once post-search distress was accounted for suggests that online searching may be one route by which health anxiety is turned into healthcare use. However, these results must not be over-interpreted. In cross-sectional data, mediation tests are tests of a hypothesised ordering, not of causation, and they can be seriously biased estimates of longitudinal processes {cite("maxwell2007")}. Our reverse-direction model (searching → distress → health anxiety) was equally significant, which is exactly what one would expect if the relationship is bidirectional, but also what one would expect if both constructs simply share variance. In addition, the mediators and outcomes in models A–C are subscales of the same instrument, so part of their association reflects shared method variance. The appropriate conclusion is that our data are *compatible with*, but cannot *demonstrate*, a cycle in which health anxiety and cyberchondria fuel each other.')
H('4.4 The role of AI chatbots', 2)
P(f'The dominance of AI chatbots ({pct(ai_n)}) is a new finding and contrasts with a 2024 US survey in which about one in five respondents reported using ChatGPT for health information {cite("ayo2024")}. Our sample was young and highly educated, which partly explains the difference: AI use fell from 96% in those aged 18–19 to 72% in those aged 40 and over. AI users had higher cyberchondria, and AI use remained an independent predictor after adjustment for health anxiety, age, occupation and other tools. The effect was concentrated in excessiveness and reassurance seeking. Two interpretations are possible. Because LLM chatbots answer in fluent, personalised and empathetic language {cite("thiru2023","ayers2023")}, they may invite repeated, conversational querying (asking "one more question"), which is the behavioural core of the excessiveness dimension. Alternatively, people who already search more heavily may simply have adopted every available tool, so that AI use marks intensity rather than causes it. The comparison group of non-users was small (n = 34) and older, the association disappeared when the 32 respondents with inconsistent answers about AI use were classed as non-users, and AI use did not modify the strength of the health anxiety–cyberchondria link. On current evidence, AI chatbots appear to be a new channel for, rather than a qualitatively different driver of, cyberchondria, which is consistent with observations that health-anxious people are over-represented among users of digital symptom checkers {cite("wetzel2024")}.')
P('Trust in online and AI information was mostly moderate and was associated with excessive searching but not with distress. Higher trust was also associated with lower odds of consulting several doctors, suggesting that trusted AI answers may substitute for some "doctor shopping". This is potentially beneficial for health services, but it is also where false reassurance becomes a risk.')
H('4.5 Pre-consultation behaviour and the doctor–patient relationship', 2)
P(f'The most common response after searching, delaying or skipping care because the information was reassuring (35%), is a double-edged finding. For self-limiting symptoms it may save unnecessary visits; for serious ones it represents a risk of delayed diagnosis. Notably, this behaviour was not related to health anxiety or cyberchondria, meaning it occurs across the spectrum and cannot be dismissed as a feature of anxious users. Conversely, cyberchondria was associated with urgent visits and emergency attendance and with fear-driven delay, mirroring evidence that health anxiety increases healthcare utilisation {cite("barsky2001")}. Thus online searching appears to push different people in opposite directions: towards complacency in some and towards over-use of services in others.')
P(f'The disclosure findings are clinically important. Almost a third of participants never told their doctor about their searches, for fear of being judged or dismissed, and those with higher health anxiety were more likely to conceal searching, whereas those with higher cyberchondria were more likely to disclose it. One interpretation is that health-anxious patients anticipate being labelled as "worried well" and hide their searching, while patients for whom searching is a dominant behaviour bring their findings to the consultation as part of reassurance seeking. Concealment matters because online information discussed openly can strengthen the doctor–patient relationship, whereas undisclosed information is more likely to undermine it {cite("tan2017")}. One in three participants reported that searching had complicated their relationship with or reduced their trust in physicians, and this was more likely at higher levels of cyberchondria.')
H('4.6 Medical and health-science students', 2)
P(f'Medical students had higher unadjusted health anxiety than non-health employees, but this difference disappeared after adjustment, largely reflecting their younger age; age was an independent predictor of clinically relevant health anxiety. The SHAI–CSS correlation tended to be stronger among medical students. This is compatible with the idea of "medical student syndrome", but evidence for that concept is weak {cite("waterman2014")}, and our data suggest that medical knowledge does not protect against cyberchondria (medical students\' CSS scores were identical to those of other participants).')
H('4.7 Comparison with regional studies', 2)
P(f'Our findings are in line with emerging Arab-region data. In Jeddah, cyberchondria was common and associated with smartphone addiction {cite("elzayat2023")}; among Saudi primary-care patients, cyberchondria had a positive effect on health anxiety in structural models {cite("pasayan2025")}; and the Arabic CSS-12 showed a structure similar to ours in Lebanon {cite("hallit2023")}. The good fit of the four-factor CSS-12 and the comparable reliability across Arabic and English responders add to the evidence that the instrument functions well in Arabic-speaking populations.')
H('4.8 Strengths and limitations', 2)
P(f'Strengths of this study include the use of two established instruments with good reliability and confirmed structure; a bilingual questionnaire with empirical checks of language equivalence; the separate assessment of AI chatbots; complete data; correction for multiple testing; robust regression, bootstrap inference and several sensitivity analyses; and the linking of psychological measures to concrete behaviours relevant to clinical practice.')
P('The limitations are substantial and should be weighed carefully:')
for b in ['**Design and causality:** the cross-sectional design cannot establish temporal order; mediation results describe statistical pathways, not causal mechanisms, and may be biased relative to longitudinal effects.',
          f'**Sampling and generalisability:** an online convenience/snowball sample is subject to self-selection. The sample was young, highly educated and over-represented medical students ({pct(medst)}) and people with health backgrounds ({pct(hb)}), so prevalence estimates and the frequency of AI use cannot be generalised to the adult population of Baghdad. Residence in Baghdad was assured by the distribution channels rather than verified by a question, and non-Internet users, older adults and less-educated people are under-represented by design.',
          '**Measurement:** the Arabic wording was produced by the research team without a formal forward–backward translation and cognitive pre-testing; equivalence was supported empirically (similar reliability, congruent loadings and equal means across language groups) but not established by formal measurement-invariance testing. "Language relied on" was self-selected, not randomised. The SHAI-14 showed only moderate fit under maximum-likelihood estimation. The SHAI cut-off was derived in Sweden, and the CSS-12 has no validated clinical cut-off, so the "high cyberchondria" group is sample-relative.',
          '**Single-item and self-designed measures:** trust, motivations, behaviours, disclosure and relationship were assessed with single, non-validated items, some allowing multiple answers; the questionnaire asked for "most frequent" behaviours rather than their frequency.',
          '**Self-report bias:** all measures were self-reported and may be affected by recall and social desirability bias; no clinical data or diagnoses were available, and chronic disease was self-reported.',
          '**Multiple testing and power:** despite FDR correction, the large number of exploratory tests means that some isolated associations may be false positives; several behavioural associations were nominally significant only. Some subgroups (non-AI users, n = 34; healthcare workers, n = 15; older age groups) were small, and the logistic model for clinical health anxiety had about five events per predictor.',
          '**Classification of AI use:** 32 respondents named an AI chatbot without ticking the AI option. They were classed as users, and the independent association between AI use and cyberchondria depended on this decision, so it should be regarded as tentative.',
          '**Data quality:** an online survey cannot exclude careless responding entirely; nine straight-line responders were identified, although their exclusion did not change the results.']:
    bullet(b)

# ------------------------------------------------------------------ CONCLUSIONS
H('5. Conclusions and Recommendations')
H('5.1 Conclusions', 2)
for b in [f'AI chatbots, mainly ChatGPT, have become the leading source of pre-consultation health information among digitally active adults in Baghdad, used by {pct(ai_n)} of participants, ahead of search engines.',
          f'About one in five participants ({prev["HA_clin"]:.0f}%) had health anxiety above a clinically relevant threshold, and younger people and those searching for instant reassurance were most affected.',
          f'Health anxiety and cyberchondria are moderately and robustly associated (r = {sc["pearson_r"]:.2f}); health anxiety is the strongest independent predictor of cyberchondria, especially of excessive searching and post-search distress.',
          'The pattern of associations is consistent with a reinforcing cycle (health anxiety → excessive searching → distress → reassurance-seeking consultations), although the cross-sectional design cannot prove its direction.',
          'AI-chatbot use is associated with higher cyberchondria, mainly excessive searching and reassurance seeking, although this depends on how inconsistent answers are classified; AI use does not change the strength of the health anxiety–cyberchondria link.',
          'Online searching shapes the consultation: one-third of participants delay care because they feel reassured, one-third never disclose their searches to their doctor, and one-third report a negative effect on their relationship with or trust in physicians. Higher cyberchondria is linked with urgent care-seeking and a more strained relationship.']:
    numbered(b)
H('5.2 Recommendations', 2)
P('**For clinicians and medical education:**', indent=False)
for b in ['Routinely and non-judgementally ask patients what they have searched for or asked an AI chatbot, to bring hidden information into the consultation and reduce concealment.',
          'Recognise repeated searching, post-search distress and repeated reassurance seeking as possible markers of health anxiety, and screen with a brief tool such as the SHAI where appropriate; offer cognitive-behavioural strategies rather than repeated reassurance or investigations.',
          'Include digital health communication, AI literacy and the management of health anxiety in undergraduate medical curricula, and support medical students who may themselves be vulnerable.']:
    bullet(b)
P('**For the Ministry of Health and public-health bodies:**', indent=False)
for b in ['Develop campaigns on digital and AI health literacy that explain both the uses and the limits of online information, especially false reassurance that leads to delayed care.',
          'Provide or endorse reliable Arabic-language health information sources and triage services so that people searching for "whether I need a doctor" are directed to safe advice.']:
    bullet(b)
P('**For developers of AI tools:**', indent=False)
for b in ['Design health responses that recognise signs of health anxiety (e.g. repeated questions about the same symptom), avoid escalating rare diagnoses without context, and encourage appropriate professional consultation, in Arabic as well as English.']:
    bullet(b)
P('**For future research:**', indent=False)
for b in ['Conduct longitudinal or experience-sampling studies to test whether health anxiety and cyberchondria reinforce each other over time, and experimental studies comparing AI chatbot and search-engine answers on post-search anxiety.',
          'Use probability sampling of the general Baghdad population, including older and less-educated adults, and validate the Arabic SHAI-14 and CSS-12 formally (forward–backward translation, measurement invariance and clinical cut-offs).',
          'Link self-reported behaviours to objective healthcare utilisation and clinical outcomes, especially delayed diagnosis following false reassurance.']:
    bullet(b)

# ------------------------------------------------------------------ REFERENCES
H('6. References')
for i, k in enumerate(order, 1):
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.8)
    p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(4); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f'{i}.\t{REFS[k]}'); r.font.size = Pt(12)
unused = [k for k in REFS if k not in order]

# ------------------------------------------------------------------ APPENDIX
page_break()
H('Appendices')
H('Appendix A. Item-by-item correlations between the SHAI-14 and CSS-12', 2)
P(f'Item-by-item Spearman correlations between all 14 SHAI-14 items and all 12 CSS-12 items (168 pairs). Correction for multiple testing was applied within this family of 168 tests: {ib["sig_fdr"]} pairs were significant after Benjamini–Hochberg FDR correction and {ib["sig_holm"]} after Holm–Bonferroni correction. Full values are given in the supplementary workbook (sheet T28_item_by_item).', indent=False)
figure(f'{FIG}/fig9_item_by_item.png', 'Item-by-item correlations between SHAI-14 and CSS-12 items (Spearman ρ). * significant after FDR correction across the 168 pairs; ** also significant after Holm–Bonferroni correction.', width=16)

page_break()
H('Appendix B. Study questionnaire', 2)
P('The bilingual (English–Arabic) questionnaire as administered through Google Forms is reproduced below (41 questions).', indent=False)
import glob
from PIL import Image as _Im
QIMGS = sorted(glob.glob(f'{OUT}/questionnaire/c-*.png'))   # pages cropped to content; they flow onto pages naturally
for qi in QIMGS:
    w_px, h_px = _Im.open(qi).size
    w_cm = min(12.5, 12.5 * w_px / 760)                     # same scale for every page; two form pages fit per A4 page
    if w_cm * h_px / w_px > 23.5: w_cm = 23.5 * w_px / h_px  # never taller than the page body
    pq = doc.add_paragraph(); pq.alignment = WD_ALIGN_PARAGRAPH.CENTER; pq.paragraph_format.space_after = Pt(4)
    pq.paragraph_format.line_spacing = 1.0
    pq.add_run().add_picture(qi, width=Cm(w_cm))

# ------------------------------------------------------------------ fill lists of tables/figures
PAGES = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else {}
def fld(kind_):
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind_); return e
def fill_list(anchor, items, kind):
    # Word TOC field over the captions (refreshes on open); cached result carries page numbers from the render pass
    prev_el = anchor
    for i, t_ in enumerate(items, 1):
        p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(2.0); p.paragraph_format.first_line_indent = Cm(-2.0)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(2.0))
        from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
        p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        if i == 1:
            r0 = p.add_run(); r0._r.append(fld('begin'))
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = f'TOC \\h \\z \\c "{kind}"'; r0._r.append(it); r0._r.append(fld('separate'))
        r = p.add_run(f'{kind} {i}.\t'); r.bold = True; r.font.size = Pt(13)
        r2 = p.add_run(t_); r2.font.size = Pt(13)
        r3 = p.add_run(f'\t{PAGES.get(f"{kind} {i}", "")}'); r3.font.size = Pt(13)
        if i == len(items): r3._r.append(fld('end'))
        prev_el._p.addnext(p._p); prev_el = p
fill_list(LT_anchor, tables_list, 'Table'); fill_list(LF_anchor, figures_list, 'Figure')

# table of contents: Word TOC field over Heading 1-2 (refreshes on open); cached entries carry rendered page numbers
start = next(i for i, (l, t) in enumerate(HEADS) if t.startswith('Summary'))
TOC = HEADS[start:]
json.dump(TOC, open(f'{OUT}/toc_entries.json', 'w'))
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
prev_el = TOC_anchor
for i, (lvl, t_) in enumerate(TOC):
    p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(3 if lvl > 1 else 5)
    p.paragraph_format.left_indent = Cm(0 if lvl == 1 else 0.8)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    if i == 0:
        r0 = p.add_run(); r0._r.append(fld('begin'))
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = 'TOC \\o "1-2" \\h \\z \\u'; r0._r.append(it); r0._r.append(fld('separate'))
    r = p.add_run(t_); r.font.size = Pt(13 if lvl == 1 else 12); r.bold = (lvl == 1)
    r3 = p.add_run(f'\t{PAGES.get("H::" + t_, "")}'); r3.font.size = Pt(13 if lvl == 1 else 12)
    if i == len(TOC) - 1: r3._r.append(fld('end'))
    prev_el._p.addnext(p._p); prev_el = p

# ------------------------------------------------------------------ page numbering
def page_number_footer(section, fmt, start, title_page=False):
    sectPr = section._sectPr
    pg = OxmlElement('w:pgNumType'); pg.set(qn('w:fmt'), fmt); pg.set(qn('w:start'), str(start)); sectPr.append(pg)
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(); r.font.name = FONT; r.font.size = Pt(12); add_field(r, 'PAGE')
    if title_page:
        section.different_first_page_header_footer = True
s1, s2 = doc.sections[0], doc.sections[1]
page_number_footer(s1, 'lowerRoman', 1, title_page=True)
page_number_footer(s2, 'decimal', 1)
for s in doc.sections:
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    s.left_margin = Cm(3.0); s.right_margin = Cm(2.5); s.top_margin = Cm(2.5); s.bottom_margin = Cm(2.5)
# ask Word to refresh fields on open
settings = doc.settings.element; uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true'); settings.append(uf)
doc.save(DOCX)
print('saved', DOCX, '| refs cited:', len(order), '| unused refs:', unused, '| tables:', len(tables_list), '| figures:', len(figures_list))
